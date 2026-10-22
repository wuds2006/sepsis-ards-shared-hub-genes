#!/usr/bin/env python3
"""Step 73 —— Response_to_Interdisciplinary_Reviewer.md -> .docx
轻量 markdown 转换: 标题/粗体/行内代码/列表/表格/引用块; Times New Roman 11pt."""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "/Users/mac/WorkBuddy/2026-08-23-21-32-45/ARDS_Sepsis_reanalysis/AJRCCM_submission/08_audit/Response_to_Interdisciplinary_Reviewer.md"
DST = SRC.replace(".md", ".docx")

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(11)


def add_runs(p, text):
    """处理 **bold** 与 *italic* 行内标记"""
    tokens = re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.font.name = "Courier New"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            p.add_run(tok)


def add_table(rows):
    ncol = len(rows[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, cell_text in enumerate(row):
            p = cells[j].paragraphs[0]
            add_runs(p, cell_text.strip())
            if i == 0:
                for r in p.runs:
                    r.bold = True


lines = open(SRC).read().splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    s = line.strip()
    if not s:
        i += 1
        continue
    # markdown 表格
    if s.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1].strip()):
        tbl = [s]
        i += 2
        while i < len(lines) and lines[i].strip().startswith("|"):
            tbl.append(lines[i].strip()); i += 1
        rows = [[c for c in r.strip("|").split("|")] for r in tbl]
        add_table(rows)
        doc.add_paragraph("")
        continue
    m = re.match(r"^(#{1,4})\s+(.*)", s)
    if m:
        level = min(len(m.group(1)), 3)
        h = doc.add_heading("", level=level)
        add_runs(h, m.group(2))
        for r in h.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
    elif s.startswith("> "):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(24)
        add_runs(p, s[2:])
    elif re.match(r"^[-*]\s+", s):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, re.sub(r"^[-*]\s+", "", s))
    elif re.match(r"^\d+\.\s+", s):
        p = doc.add_paragraph(style="List Number")
        add_runs(p, re.sub(r"^\d+\.\s+", "", s))
    else:
        p = doc.add_paragraph()
        add_runs(p, s)
    i += 1

doc.save(DST)
print("Saved", DST)
