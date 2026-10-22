#!/usr/bin/env python3
# step95: convert revised manuscript md -> docx (Times New Roman 12, gene italics)
import re
from docx import Document
from docx.shared import Pt

SRC = '/Users/mac/WorkBuddy/2026-08-23-21-32-45/ARDS_Sepsis_reanalysis/AJRCCM_submission/09_review_deAI_2026-09-25/02_Manuscript_Revised.md'
DST = '/Users/mac/WorkBuddy/2026-08-23-21-32-45/ARDS_Sepsis_reanalysis/AJRCCM_submission/09_review_deAI_2026-09-25/manuscript_revised.docx'

HEADS = {'ABSTRACT', 'INTRODUCTION', 'METHODS', 'RESULTS', 'DISCUSSION',
         'CONCLUSION', 'REFERENCES', 'DATA AVAILABILITY'}

SUBHEADS = {
    # Results subheads after the Path-A restructure of 2026-10-16
    'Identification of Shared Candidate Genes',
    'GO and KEGG Enrichment of Shared Candidate Genes',
    'Machine Learning Identifies Six Hub Genes',
    'Expression Patterns and Enrichment of the Hub Genes',
    'Disease Prediction Model Construction',
    'External Confirmation of the Six-Gene Signature',
    'Prognostic Performance in Septic Shock',
    'Mendelian Randomization Under Limited Power',
    'Immune Microenvironment: Myeloid Expansion and Complement Activation',
    'Single-Cell Localization of the Hub Genes',
    'Protein-Level Validation and Therapeutic Relevance',
}

def add_inline(paragraph, text):
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            paragraph.add_run(tok[2:-2])  # bold stripped per de-AI constraint
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1]); r.italic = True
        else:
            paragraph.add_run(tok)

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

lines = open(SRC, encoding='utf-8').read().split('\n')
first_title_done = False
for line in lines:
    s = line.rstrip()
    if not s.strip() or s.strip() == '---':
        continue
    if s in HEADS:
        h = doc.add_heading(level=1); add_inline(h, s)
    elif not first_title_done:
        h = doc.add_heading(level=1); add_inline(h, s)
        first_title_done = True
    elif s.startswith('Running head:') or s.startswith('Article type:'):
        p = doc.add_paragraph(); add_inline(p, s)
    elif re.match(r'^\d+\.\s', s):
        p = doc.add_paragraph(); add_inline(p, s)  # references as plain numbered text
    elif re.match(r'^(Rationale|Objectives|Methods|Measurements and Main Results|Conclusions|Keywords):', s):
        p = doc.add_paragraph(); add_inline(p, s)
    elif s in SUBHEADS:
        h = doc.add_heading(level=2); add_inline(h, s)
    else:
        p = doc.add_paragraph(); add_inline(p, s)

doc.save(DST)
print('saved', DST)
