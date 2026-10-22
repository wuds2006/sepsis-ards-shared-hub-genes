#!/usr/bin/env python
"""Build supplementary Tables S1-S8 as xlsx workbooks into the submission package.
Sources follow Supplementary_Materials.md legend source lines.
"""
import csv, os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

ROOT = "/Users/mac/WorkBuddy/2026-08-23-21-32-45/ARDS_Sepsis_reanalysis"
OUT = os.path.join(ROOT, "AJRCCM_submission/05_supplementary/Supplementary_Figures_Tables/Tables")
os.makedirs(OUT, exist_ok=True)

HDR_FONT = Font(bold=True)

def add_sheet(wb, title, rows, max_rows=None):
    ws = wb.create_sheet(title[:31])
    if max_rows is not None:
        rows = rows[:max_rows]
    for r, row in enumerate(rows):
        ws.append(row)
        if r == 0:
            for cell in ws[1]:
                cell.font = HDR_FONT
    return ws

def read_csv(path, delimiter=","):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f, delimiter=delimiter))

def build(name, sheets, note=None):
    wb = Workbook()
    wb.remove(wb.active)
    if note:
        ws = wb.create_sheet("README")
        for line in note:
            ws.append([line])
    for title, rows in sheets:
        add_sheet(wb, title, rows)
    path = os.path.join(OUT, name)
    wb.save(path)
    print("saved", path)

P = lambda *a: os.path.join(ROOT, *a)

# ---- Table S1: full DEG results ----
build("TableS1_DEG_full.xlsx", [
    ("GSE76293_ARDS_vs_Control", read_csv(P("01_DEG", "GSE76293_DEG_all.csv"))),
    ("GSE69528_Sepsis_vs_Control", read_csv(P("01_DEG", "GSE69528_DEG_all.csv"))),
])

# ---- Table S2: candidate gene lists ----
def gene_list(path, colname):
    with open(path) as f:
        g = [l.strip() for l in f if l.strip()]
    return [[colname]] + [[x] for x in g]

build("TableS2_Shared_candidate_genes.xlsx", [
    ("Shared_CGs_179", gene_list(P("03_intersection", "Shared_CGs.txt"), "gene")),
    ("ARDS_CGs", gene_list(P("03_intersection", "ARDS_CGs.txt"), "gene")),
    ("Sepsis_CGs", gene_list(P("03_intersection", "Sepsis_CGs.txt"), "gene")),
])

# ---- Table S3: MR + SMR/HEIDI + coloc ----
hub_ensg = {"CDK5RAP2": "ENSG00000136861", "CD177": "ENSG00000204936",
            "GRB10": "ENSG00000106070", "UPP1": "ENSG00000183696",
            "VNN1": "ENSG00000112299"}
smr_rows = [["gene", "outcome", "probeID", "topSNP", "b_SMR", "se_SMR",
             "p_SMR", "p_HEIDI", "nsnp_HEIDI"]]
for gene, ensg in hub_ensg.items():
    for outcome in ("sepsis", "ards"):
        fn = P("04_MR_SMR", f"smr_{gene}_{outcome}_heidi.smr")
        with open(fn) as f:
            for r in csv.DictReader(f, delimiter="\t"):
                if r["probeID"] == ensg or r["Gene"] == ensg:
                    smr_rows.append([gene, outcome, r["probeID"], r["topSNP"],
                                     r["b_SMR"], r["se_SMR"], r["p_SMR"],
                                     r["p_HEIDI"], r["nsnp_HEIDI"]])

build("TableS3_MR_SMR_coloc.xlsx", [
    ("MR_estimates", read_csv(P("04_MR_SMR", "MR_sepsis4980_ARDS_results.csv"))),
    ("MR_heterogeneity", read_csv(P("04_MR_SMR", "MR_sepsis4980_heterogeneity.csv"))),
    ("MR_pleiotropy", read_csv(P("04_MR_SMR", "MR_sepsis4980_pleiotropy.csv"))),
    ("MR_Steiger", read_csv(P("04_MR_SMR", "MR_sepsis4980_steiger.csv"))),
    ("SMR_HEIDI_hub_genes", smr_rows),
    ("coloc_results", read_csv(P("04_MR_SMR", "coloc_results.csv"))),
], note=["Supplementary Table S3. Mendelian randomization (sepsis ieu-b-4980 -> ARDS FinnGen J10_ARDS),",
         "SMR-HEIDI (eQTLGen blood cis-eQTL) and coloc results for the six hub genes.",
         "EXOSC4 had no blood cis-eQTL and was not testable by SMR."])

# ---- Table S4: immune infiltration ----
build("TableS4_Immune_infiltration.xlsx", [
    ("GSE76293_ssGSEA_diff", read_csv(P("07_hub_immune", "GSE76293_immune_diff.csv"))),
    ("GSE69528_ssGSEA_diff", read_csv(P("07_hub_immune", "GSE69528_immune_diff.csv"))),
    ("GSE76293_hub_immune_cor", read_csv(P("07_hub_immune", "GSE76293_hub_immune_cor.csv"))),
    ("GSE69528_hub_immune_cor", read_csv(P("07_hub_immune", "GSE69528_hub_immune_cor.csv"))),
])

# ---- Table S5: GSVA ----
build("TableS5_GSVA_537_pathways.xlsx", [
    ("GSE76293_GSVA_limma", read_csv(P("07_hub_immune", "gsva_top20", "GSE76293_GSVA_limma_all.csv"))),
    ("GSE69528_GSVA_limma", read_csv(P("07_hub_immune", "gsva_top20", "GSE69528_GSVA_limma_all.csv"))),
])

# ---- Table S6: docking ----
build("TableS6_Molecular_docking.xlsx", [
    ("docking_site_results", read_csv(P("09_docking_drug", "docking_site_results.csv"))),
])

# ---- Table S7: AUCell ----
build("TableS7_AUCell_stats.xlsx", [
    ("GSE167363_stats", read_csv(P("08_singlecell", "GSE167363", "aucell", "GSE167363_AUCell_stats.csv"))),
    ("GSE167363_genesets", read_csv(P("08_singlecell", "GSE167363", "aucell", "GSE167363_AUCell_genesets.csv"))),
    ("GSE151263_stats", read_csv(P("08_singlecell", "aucell", "GSE151263_AUCell_stats.csv"))),
    ("GSE151263_genesets", read_csv(P("08_singlecell", "aucell", "GSE151263_AUCell_genesets.csv"))),
])

# ---- Table S8: virtual knockout ----
KO = P("08_singlecell", "GSE167363", "knockout")
build("TableS8_Virtual_knockout.xlsx", [
    ("KO_feasibility", read_csv(os.path.join(KO, "KO_feasibility.csv"))),
    ("KO_UPP1_diffRegulation", read_csv(os.path.join(KO, "KO_UPP1_diffRegulation.csv"))),
    ("KO_CDK5RAP2_diffRegulation", read_csv(os.path.join(KO, "KO_CDK5RAP2_diffRegulation.csv"))),
    ("KO_EXOSC4_diffRegulation", read_csv(os.path.join(KO, "KO_EXOSC4_diffRegulation.csv"))),
    ("KO_VNN1_diffRegulation", read_csv(os.path.join(KO, "KO_VNN1_diffRegulation.csv"))),
    ("KO_hub_cross_effects", read_csv(os.path.join(KO, "KO_hub_cross_effects.csv"))),
    ("CellOracle_TF_KO_hub_effect", read_csv(P("08_singlecell", "GSE167363", "celloracle", "CellOracle_TF_KO_hub_effect.csv"))),
])

print("DONE")
